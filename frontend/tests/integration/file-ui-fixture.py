"""Synthetic assets and exact read-only controls for the real file UI packet."""
import json
from hashlib import sha256
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
ASSETS = ROOT / '.local/resume-file-assets'


def main():
    if sys.argv[1] == 'pdf':
        from PIL import Image, ImageDraw, ImageFont
        import pymupdf
        ASSETS.mkdir(exist_ok=False)
        image = Image.new('RGB', (1600, 700), 'white')
        ImageDraw.Draw(image).multiline_text((80, 100), 'САПФИР-742 / SAPPHIRE-742\nСтраница двадцать пять',
            font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 64), fill='black', spacing=45)
        image.save(ASSETS / 'scan.png')
        document = pymupdf.open()
        for number in range(1, 26):
            page = document.new_page()
            if number == 25:
                page.insert_image(pymupdf.Rect(20, 100, 575, 343), filename=str(ASSETS / 'scan.png'))
            else:
                page.insert_text((50, 100), f'Digital page {number}: controlled UI document.')
        document.save(ASSETS / 'scanned-page-25.pdf')
        assert len(document) == 25 and not document[24].get_text() and document[24].get_images()
        (ASSETS / 'statement.txt').write_text('Контрольный код САПФИР-742 / SAPPHIRE-742. Проверка русского текста.', encoding='utf-8')
        print('Synthetic TXT and 25-page mixed PDF prepared')
        return
    if sys.argv[1] == 'word':
        from docx import Document
        document = Document()
        document.add_paragraph('Контрольный код САПФИР-742. Изображение внутри DOCX не распознаётся.')
        document.add_picture(str(ASSETS / 'scan.png'))
        document.save(ASSETS / 'partial-image.docx')
        assert len(document.paragraphs) == 2 and len(document.inline_shapes) == 1
        print('Synthetic Word document with text and image prepared')
        return
    import psycopg
    from psycopg.rows import dict_row
    from backend.files import read_chunks, storage_path
    inputs = json.load(sys.stdin)
    info = json.loads((ROOT / '.local/p12-ui-connection.json').read_text(encoding='utf-8'))
    assert info['dbname'] == 'deanery_p12_ui_92004917'
    with psycopg.connect(**info, row_factory=dict_row) as db:
        db.execute('SET TRANSACTION READ ONLY')
        assert db.execute('SELECT version_num FROM public.alembic_version').fetchone()['version_num'] == '0039_scope_check_latency'
        assert db.execute("SELECT count(*) AS n FROM agent_backend.agent_run WHERE finished_at IS NULL").fetchone()['n'] == 0
        assert db.execute("SELECT count(*) AS n FROM agent_backend.job WHERE status IN ('queued','running')").fetchone()['n'] == 0
        assert db.execute('SELECT count(*) AS n FROM agent_backend.auth_session WHERE revoked_at IS NULL').fetchone()['n'] == 5
        if sys.argv[1] == 'inventory':
            assert db.execute('SELECT count(*) AS n FROM agent_backend.file').fetchone()['n'] == 0
            print(json.dumps({'head': '0039_scope_check_latency', 'files': 0, 'active_runs_jobs': 0, 'historical_sessions': 5}))
            return
        assert sys.argv[1] == 'verify' and len(inputs['uploads']) == 3
        rows = []
        for item in inputs['uploads']:
            descriptor = db.execute('SELECT agent_backend.file_descriptor(%s) AS result', (item['version_id'],)).fetchone()['result']
            original = ASSETS / item['filename']
            digest = sha256(original.read_bytes()).hexdigest()
            assert descriptor['content_hash'] == digest
            assert sha256(storage_path(descriptor['storage_key'], ROOT / '.local/p12-ui-files').read_bytes()).hexdigest() == digest
            chunks = read_chunks(descriptor, root=ROOT / '.local/p12-ui-files', **({'page': 25} if original.suffix == '.pdf' else {}))['chunks']
            assert chunks and any('742' in chunk['text'] for chunk in chunks)
            coverage = descriptor['extraction']['coverage']
            if original.suffix == '.pdf':
                assert coverage['total_pages'] == coverage['attempted_pages'] == 25
                assert coverage['ocr_pages'] == 1 and coverage['failed_pages'] == coverage['unprocessed_pages'] == 0
                assert all(chunk['page'] == 25 and chunk['source'] == 'ocr' for chunk in chunks)
            elif original.suffix == '.docx':
                assert descriptor['extraction_state'] == 'partial'
                assert 'docx_images_not_extracted' in descriptor['extraction']['warnings'] and not coverage['complete']
            else:
                assert descriptor['extraction_state'] == 'ready' and coverage['complete']
            rows.append({'filename': item['filename'], 'hash_matches': True, 'extraction_state': descriptor['extraction_state'], 'coverage': coverage, 'selected_chunks': len(chunks)})
        print(json.dumps({'files': rows, 'provider_calls': 0, 'historical_sessions_preserved': 5}))


if __name__ == '__main__':
    assert Path.cwd().resolve() == ROOT
    main()
